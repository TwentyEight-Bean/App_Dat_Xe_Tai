import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    COLOR_PRIMARY = RGBColor(15, 23, 42)      # Slate 900
    COLOR_BRAND = RGBColor(2, 132, 199)       # Sky 600
    COLOR_ACCENT = RGBColor(217, 119, 6)      # Amber 600
    COLOR_TEXT = RGBColor(51, 65, 85)         # Slate 700

    # ==========================================
    # HEADER / TITLE
    # ==========================================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = title_p.add_run("HỆ THỐNG ĐẶT XE TẢI CÔNG NGHỆ FASTDROP\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_BRAND

    run_title = title_p.add_run("KẾ HOẠCH & KỊCH BẢN KIỂM THỬ ỨNG DỤNG\n(TEST PLAN & USE CASE SCENARIOS)")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(18)
    run_desc = sub_p.add_run("Vai trò: Khách hàng (Customer) & Đối tác Tài xế (Driver) | Phiên bản: 1.0.0 (Production Ready)")
    run_desc.font.name = "Arial"
    run_desc.font.size = Pt(10)
    run_desc.font.italic = True
    run_desc.font.color.rgb = RGBColor(100, 116, 139)

    # ==========================================
    # PHẦN 1: THÔNG TIN MÔI TRƯỜNG & TÀI KHOẢN MẪU
    # ==========================================
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. MÔI TRƯỜNG & DỮ LIỆU KIỂM THỬ SẴN CÓ")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(13)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_BRAND

    t_info = doc.add_table(rows=6, cols=3)
    t_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_info.autofit = False

    info_headers = ["Hạng mục", "Thông số / Địa chỉ truy cập", "Ghi chú hướng dẫn"]
    info_data = [
        ["Đường dẫn Web (Máy tính)", "http://localhost:8443", "Mở trực tiếp trên Chrome / Edge"],
        ["Đường dẫn Mobile (Điện thoại)", "http://172.16.149.59:8443", "Bắt chung Wi-Fi với máy tính"],
        ["Tài khoản Khách hàng mẫu", "SĐT: 0988888888 | PIN: 123456", "Đã có sẵn ví và lịch sử đặt xe"],
        ["Tài khoản Tài xế mẫu 1", "SĐT: 0908815624 | PIN: 123456", "Tài xế Nguyễn Văn Minh (Xe 1 Tấn - 51C-888.88)"],
        ["Tài khoản Tài xế mẫu 2", "SĐT: 0903124486 | PIN: 123456", "Tài xế Trần Quốc Hùng (Xe 2 Tấn - 51D-999.99)"],
    ]

    # Format header
    hdr_cells = t_info.rows[0].cells
    for i, title in enumerate(info_headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "0F172A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Format data rows
    for row_idx, row in enumerate(info_data):
        cells = t_info.rows[row_idx + 1].cells
        for col_idx, text in enumerate(row):
            cells[col_idx].text = text
            bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cells[col_idx], bg_color)
            set_cell_margins(cells[col_idx], top=80, bottom=80, left=100, right=100)
            p = cells[col_idx].paragraphs[0]
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_TEXT

    # Set column widths
    col_widths = [Inches(1.8), Inches(2.6), Inches(2.6)]
    for row in t_info.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # ==========================================
    # PHẦN 2: KỊCH BẢN DÀNH CHO KHÁCH HÀNG (CUSTOMER)
    # ==========================================
    doc.add_paragraph().paragraph_format.space_before = Pt(14)
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. KỊCH BẢN KIỂM THỬ: VAI TRÒ KHÁCH HÀNG (CUSTOMER)")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(13)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_BRAND

    customer_cases = [
        {
            "code": "TC-KH-01",
            "name": "Đăng ký tài khoản mới bắt buộc OTP (Chống số ảo)",
            "steps": "1. Mở modal Đăng nhập -> Bấm 'Đăng ký mới'\n2. Nhập Họ tên, SĐT mới (VD: 097xxxxxxx), tạo Mã PIN (VD: 654321)\n3. Bấm 'Xác minh SĐT qua OTP & Tiếp tục'\n4. Nhập sai OTP (VD: 000000) -> Xem hệ thống báo lỗi\n5. Nhập đúng mã OTP (hoặc bấm 'Điền nhanh' 123456) -> Hoàn tất",
            "expected": "• Hệ thống gửi SMS OTP 6 số, hiển thị đếm ngược 60s.\n• Nhập sai OTP bị chặn và đếm số lần còn lại.\n• Nhập đúng OTP -> Tài khoản được tạo thành công, tự động lưu session và đăng nhập vào màn hình chính."
        },
        {
            "code": "TC-KH-02",
            "name": "Đăng nhập bằng Mã PIN & Đăng nhập nhanh OTP",
            "steps": "1. Mở form Đăng nhập, nhập SĐT: 0988888888\n2. Nhập sai Mã PIN (VD: 111111) -> Bấm Đăng nhập\n3. Nhập đúng Mã PIN: 123456 -> Bấm Đăng nhập\n4. Thử đăng xuất, chọn 'Đăng nhập nhanh bằng OTP' -> Nhận OTP và vào app.",
            "expected": "• Nhập sai PIN: Báo lỗi đỏ rõ ràng.\n• Nhập đúng PIN: Đăng nhập thành công tức thì, hiển thị đúng họ tên 'Khách hàng FastDrop'.\n• Đăng nhập OTP: Vào app thành công mà không cần nhớ mật khẩu."
        },
        {
            "code": "TC-KH-03",
            "name": "Chọn lộ trình & Tính cước minh bạch (Interactive Map)",
            "steps": "1. Tại màn hình Trang chủ, gõ địa chỉ giao hàng hoặc chạm kéo ghim trên bản đồ.\n2. Quan sát lộ trình vẽ nối giữa Điểm đón và Điểm đến.\n3. Xem thẻ cước phí phía dưới.",
            "expected": "• Bản đồ hiển thị đường lộ trình màu xanh rõ nét.\n• Hệ thống tính chính xác số km và thời gian dự kiến (VD: 1.1 km • ~5 phút • 200.000đ).\n• Hiển thị đúng nhãn 'Giá trọn gói theo km' (Không bị hiện sai 'Vui lòng chọn điểm đón & đến')."
        },
        {
            "code": "TC-KH-04",
            "name": "Lựa chọn loại xe tải & Dịch vụ phụ phí bốc xếp",
            "steps": "1. Lần lượt bấm chọn các loại xe: Xe 500kg, Xe 1 Tấn, Xe 2 Tấn, Xe 3 Tấn.\n2. Bấm vào 'Dịch vụ bốc xếp & Phụ phí' -> Chọn 'Bốc xếp tầng trệt' (+50.000đ) hoặc 'Giao tận nơi'.\n3. Quan sát tổng giá cước.",
            "expected": "• Giá tiền nhảy số mượt mà theo từng loại xe.\n• Thêm phụ phí -> Thẻ cước đổi thành nhãn 'Đã gồm phụ phí +50.000đ' chính xác.\n• Bỏ chọn phụ phí -> Giá trở lại cước chuẩn."
        },
        {
            "code": "TC-KH-05",
            "name": "Phương thức thanh toán: Tiền mặt vs Chuyển khoản VietQR",
            "steps": "1. Tại mục Phương thức thanh toán, quan sát 2 lựa chọn: 💵 Tiền mặt và 💳 Chuyển khoản.\n2. Bấm chọn 'Chuyển khoản' -> Bấm nút 'Tìm tài xế'\n3. Quan sát Popup quét mã VietQR xuất hiện.",
            "expected": "• Chỉ có 2 nút thanh toán gọn gàng, bấm nút nào nút đó sáng viền xanh active.\n• Khi chọn Chuyển khoản: Bật Popup VietQR kèm số tiền, ngân hàng và nội dung chuyển khoản tự động."
        },
        {
            "code": "TC-KH-06",
            "name": "Khách hàng hủy chuyến xe (Cancellation Policy)",
            "steps": "1. Đặt xe ở trạng thái đang tìm tài xế -> Bấm 'Hủy chuyến'\n2. Khi tài xế đã nhận đơn hoặc đã đến điểm đón -> Thử bấm 'Hủy chuyến'.",
            "expected": "• Hủy khi chưa có xe: Hủy ngay lập tức, hoàn cọc 100%, không phạt.\n• Hủy khi tài xế đã đến điểm đón: Bật modal cảnh báo áp dụng phí bồi thường cho tài xế."
        },
        {
            "code": "TC-KH-07",
            "name": "Đánh giá sao & Nhận xét tài xế sau hoàn thành",
            "steps": "1. Sau khi chuyến xe kết thúc (COMPLETED), màn hình hiện popup Đánh giá.\n2. Chọn 5 sao, bấm chọn tag 'Thân thiện', 'Đúng giờ', gõ nhận xét -> Bấm 'Gửi đánh giá'.",
            "expected": "• Lưu điểm đánh giá thành công vào hệ thống, cập nhật điểm trung bình của tài xế."
        }
    ]

    t_cust = doc.add_table(rows=len(customer_cases) + 1, cols=4)
    t_cust.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cust.autofit = False

    cust_headers = ["Mã TC", "Tên tình huống kiểm thử", "Các bước thực hiện", "Kết quả mong đợi (Expected Result)"]
    c_cells = t_cust.rows[0].cells
    for i, title in enumerate(cust_headers):
        c_cells[i].text = title
        set_cell_background(c_cells[i], "0284C7")
        set_cell_margins(c_cells[i], top=100, bottom=100, left=120, right=120)
        p = c_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, tc in enumerate(customer_cases):
        cells = t_cust.rows[row_idx + 1].cells
        cells[0].text = tc["code"]
        cells[1].text = tc["name"]
        cells[2].text = tc["steps"]
        cells[3].text = tc["expected"]

        bg_color = "F0F9FF" if row_idx % 2 == 0 else "FFFFFF"
        for c in cells:
            set_cell_background(c, bg_color)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.color.rgb = COLOR_TEXT

    cust_col_widths = [Inches(1.0), Inches(1.8), Inches(2.2), Inches(2.0)]
    for row in t_cust.rows:
        for idx, width in enumerate(cust_col_widths):
            row.cells[idx].width = width

    # ==========================================
    # PHẦN 3: KỊCH BẢN DÀNH CHO TÀI XẾ (DRIVER)
    # ==========================================
    doc.add_paragraph().paragraph_format.space_before = Pt(14)
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("3. KỊCH BẢN KIỂM THỬ: VAI TRÒ TÀI XẾ ĐỐI TÁC (DRIVER)")
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(13)
    r_h3.font.bold = True
    r_h3.font.color.rgb = COLOR_ACCENT

    driver_cases = [
        {
            "code": "TC-TX-01",
            "name": "Đăng nhập đối tác tài xế & Bật Trực tuyến (Online)",
            "steps": "1. Chuyển sang Cổng Tài xế (bấm nút góc trên 'Tài xế').\n2. Bấm chọn nhanh tài xế 'Anh Minh (Xe 1 Tấn - 51C-888.88)'.\n3. Đăng nhập vào trang điều khiển tài xế -> Gạt nút 'Trực tuyến' sang ON.",
            "expected": "• Đăng nhập thành công, nạp thông tin xe và số dư ví.\n• Nút Online chuyển sang màu xanh lá, hệ thống sẵn sàng nhận cuốc trong bán kính 10km."
        },
        {
            "code": "TC-TX-02",
            "name": "Nổ cuốc xe mới & Nhận cuốc trong 30 giây",
            "steps": "1. Giữ tài xế ở trạng thái Online.\n2. Mở 1 tab khác đóng vai Khách hàng bấm Đặt xe.\n3. Quan sát màn hình tài xế khi có cuốc mới.\n4. Bấm nút 'NHẬN CUỐC' trước khi đồng hồ 30s đếm về 0.",
            "expected": "• Màn hình tài xế bật Popup nổ đơn kèm âm thanh chuông báo sinh động.\n• Hiển thị chi tiết: Thu nhập thực nhận, km đến điểm lấy, loại hàng, yêu cầu bốc xếp.\n• Bấm Nhận cuốc -> Khóa chuyến xe độc quyền cho tài xế, chuyển sang màn hình di chuyển."
        },
        {
            "code": "TC-TX-03",
            "name": "Cập nhật tiến trình chuyến xe (4 bước chuẩn Logistics)",
            "steps": "1. Bước 1: Tài xế di chuyển và bấm 'Đã đến điểm lấy'.\n2. Bước 2: Bấm 'Bắt đầu di chuyển giao hàng' (có tải ảnh hàng hóa).\n3. Bước 3: Di chuyển đến nơi và bấm 'Đã đến điểm giao'.\n4. Bước 4: Thu tiền (nếu tiền mặt) và bấm 'Hoàn tất chuyến xe'.",
            "expected": "• Mỗi bước cập nhật đều bắn sự kiện Realtime (Pusher) sang phía khách hàng ngay lập tức.\n• Nút chuyển trạng thái phản hồi mượt mà, không bị lag giật.\n• Hoàn tất -> Hiện màn hình tổng kết thu nhập chuyến đi."
        },
        {
            "code": "TC-TX-04",
            "name": "Chống nhận nhiều cuốc cùng lúc (1 Xe - 1 Cuốc)",
            "steps": "1. Khi tài xế đang trong chuyến đi chưa hoàn tất (IN_TRANSIT).\n2. Thử mô phỏng hoặc bấm nhận thêm một cuốc xe khác.",
            "expected": "• Hệ thống lập tức chặn lại và báo lỗi: 'Bạn đang có chuyến xe chưa hoàn thành. Không thể nhận thêm cuốc mới'."
        },
        {
            "code": "TC-TX-05",
            "name": "Quản lý Ví ký quỹ & Tự động trừ hoa hồng sàn",
            "steps": "1. Xem số dư ví trước chuyến đi.\n2. Hoàn thành chuyến xe tiền mặt 200.000đ (hoa hồng sàn 20% = 40.000đ).\n3. Kiểm tra biến động số dư trong tab Ví.",
            "expected": "• Ví tài xế tự động bị trừ đúng 40.000đ tiền hoa hồng sàn.\n• Tab Lịch sử ví ghi nhận giao dịch: 'Khấu trừ phí nền tảng cuốc xe #BK...'."
        },
        {
            "code": "TC-TX-06",
            "name": "Nạp tiền vào ví tài xế qua chuyển khoản QR",
            "steps": "1. Trong tab Ví, bấm nút 'Nạp tiền'.\n2. Chọn mệnh giá (VD: 200.000đ) -> Bấm 'Tạo mã nạp tiền'.\n3. Quét mã QR hoặc bấm xác nhận nạp tiền thử nghiệm.",
            "expected": "• Hiển thị mã QR VietQR nạp ví chính xác.\n• Sau khi nạp thành công -> Số dư ví nhảy tăng ngay lập tức."
        }
    ]

    t_drv = doc.add_table(rows=len(driver_cases) + 1, cols=4)
    t_drv.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_drv.autofit = False

    drv_headers = ["Mã TC", "Tên tình huống kiểm thử", "Các bước thực hiện", "Kết quả mong đợi (Expected Result)"]
    d_cells = t_drv.rows[0].cells
    for i, title in enumerate(drv_headers):
        d_cells[i].text = title
        set_cell_background(d_cells[i], "D97706")
        set_cell_margins(d_cells[i], top=100, bottom=100, left=120, right=120)
        p = d_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, tc in enumerate(driver_cases):
        cells = t_drv.rows[row_idx + 1].cells
        cells[0].text = tc["code"]
        cells[1].text = tc["name"]
        cells[2].text = tc["steps"]
        cells[3].text = tc["expected"]

        bg_color = "FFFBEB" if row_idx % 2 == 0 else "FFFFFF"
        for c in cells:
            set_cell_background(c, bg_color)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.color.rgb = COLOR_TEXT

    for row in t_drv.rows:
        for idx, width in enumerate(cust_col_widths):
            row.cells[idx].width = width

    # ==========================================
    # PHẦN 4: KỊCH BẢN PHỐI HỢP ĐẦU CUỐI (END-TO-END)
    # ==========================================
    doc.add_paragraph().paragraph_format.space_before = Pt(14)
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("4. KỊCH BẢN KIỂM THỬ ĐẦU CUỐI HAI BÊN (END-TO-END)")
    r_h4.font.name = "Arial"
    r_h4.font.size = Pt(13)
    r_h4.font.bold = True
    r_h4.font.color.rgb = COLOR_PRIMARY

    p_e2e_intro = doc.add_paragraph()
    p_e2e_intro.paragraph_format.space_after = Pt(8)
    r_e2e_intro = p_e2e_intro.add_run("Cách thực hiện tốt nhất: Mở 2 cửa sổ trình duyệt song song (hoặc 1 máy tính đóng vai Khách hàng, 1 điện thoại đóng vai Tài xế qua Wi-Fi http://172.16.149.59:8443).")
    r_e2e_intro.font.name = "Arial"
    r_e2e_intro.font.size = Pt(9.5)
    r_e2e_intro.font.italic = True

    e2e_steps = [
        ("Giai đoạn 1: Chuẩn bị", "Tài xế đăng nhập số 0908815624 -> Bật chế độ Trực tuyến (Online). Bản đồ khách hàng lập tức thấy icon xe tải xuất hiện gần khu vực."),
        ("Giai đoạn 2: Khách đặt xe", "Khách hàng chọn điểm giao hàng -> Chọn Xe 1 Tấn -> Chọn thanh toán Tiền mặt -> Bấm 'Tìm tài xế'."),
        ("Giai đoạn 3: Nổ cuốc & Nhận", "Màn hình tài xế đổ chuông nổ cuốc -> Tài xế bấm 'Nhận cuốc'. Màn hình khách hàng lập tức chuyển từ 'Đang tìm xe' sang 'Tài xế đã nhận chuyến' (hiện biển số 51C-888.88, sđt, tên tài xế)."),
        ("Giai đoạn 4: Chat Realtime", "Khách hàng bấm icon Chat -> Gửi tin nhắn: 'Bác tài đến cổng số 2 nhé'. Màn hình tài xế nhận ngay tin nhắn tức thì (không cần reload)."),
        ("Giai đoạn 5: Tiến trình lấy hàng", "Tài xế bấm 'Đã đến điểm lấy' -> Màn hình khách hàng đổi trạng thái 'Tài xế đã đến nơi'. Tài xế bấm 'Bắt đầu di chuyển' -> Khách hàng thấy icon xe tải lăn bánh trên bản đồ."),
        ("Giai đoạn 6: Giao hàng & Trừ ví", "Tài xế bấm 'Hoàn tất chuyến xe' -> Thu tiền mặt 200.000đ. Hệ thống tự động trừ 40.000đ tiền hoa hồng trong ví tài xế."),
        ("Giai đoạn 7: Đánh giá dịch vụ", "Màn hình khách hàng hiện thông báo chuyến xe hoàn tất -> Khách chấm 5 sao và đánh giá tài xế thân thiện.")
    ]

    t_e2e = doc.add_table(rows=len(e2e_steps) + 1, cols=3)
    t_e2e.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_e2e.autofit = False

    e2e_headers = ["Thứ tự", "Giai đoạn kiểm thử", "Diễn biến phối hợp giữa Khách & Tài xế"]
    e_cells = t_e2e.rows[0].cells
    for i, title in enumerate(e2e_headers):
        e_cells[i].text = title
        set_cell_background(e_cells[i], "1E293B")
        set_cell_margins(e_cells[i], top=100, bottom=100, left=120, right=120)
        p = e_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, (phase, detail) in enumerate(e2e_steps):
        cells = t_e2e.rows[row_idx + 1].cells
        cells[0].text = f"Bước {row_idx + 1}"
        cells[1].text = phase
        cells[2].text = detail

        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c in cells:
            set_cell_background(c, bg_color)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.color.rgb = COLOR_TEXT

    e2e_col_widths = [Inches(0.9), Inches(1.8), Inches(4.3)]
    for row in t_e2e.rows:
        for idx, width in enumerate(e2e_col_widths):
            row.cells[idx].width = width

    # ==========================================
    # PHẦN 5: CHECKLIST ĐÁNH GIÁ (PASS / FAIL)
    # ==========================================
    doc.add_paragraph().paragraph_format.space_before = Pt(14)
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    r_h5 = h5.add_run("5. BẢNG CHECKLIST ĐÁNH GIÁ NGHIỆM THU (PASS / FAIL)")
    r_h5.font.name = "Arial"
    r_h5.font.size = Pt(13)
    r_h5.font.bold = True
    r_h5.font.color.rgb = COLOR_BRAND

    checklist_items = [
        "TC-KH-01: Đăng ký bắt buộc OTP (chống số điện thoại ảo thành công)",
        "TC-KH-02: Đăng nhập SĐT + Mã PIN hoặc OTP nhanh",
        "TC-KH-03: Kéo thả bản đồ, hiển thị lộ trình và giá trọn gói theo km",
        "TC-KH-04: Chọn loại xe tải và cộng phụ phí bốc xếp chính xác",
        "TC-KH-05: Thanh toán chỉ gồm 2 phương thức: Tiền mặt & Chuyển khoản QR",
        "TC-KH-06: Khách hàng hủy cuốc an toàn theo quy định",
        "TC-KH-07: Đánh giá chất lượng chuyến xe 5 sao",
        "TC-TX-01: Tài xế đăng nhập và bật chế độ Online hoạt động",
        "TC-TX-02: Nổ cuốc rung chuông và nhận đơn trong 30 giây",
        "TC-TX-03: Cập nhật 4 trạng thái giao hàng Realtime thông suốt",
        "TC-TX-04: Chống nhận nhiều cuốc cùng lúc (1 xe - 1 cuốc độc quyền)",
        "TC-TX-05: Tự động khấu trừ hoa hồng sàn vào ví tài xế chính xác",
        "TC-TX-06: Nạp tiền ví ký quỹ tài xế qua mã VietQR thành công",
        "E2E-01: Toàn bộ chu trình tương tác Khách hàng <-> Tài xế thời gian thực"
    ]

    t_chk = doc.add_table(rows=len(checklist_items) + 1, cols=3)
    t_chk.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_chk.autofit = False

    chk_headers = ["STT", "Nội dung kiểm thử", "Kết quả kiểm tra thực tế"]
    k_cells = t_chk.rows[0].cells
    for i, title in enumerate(chk_headers):
        k_cells[i].text = title
        set_cell_background(k_cells[i], "0F172A")
        set_cell_margins(k_cells[i], top=100, bottom=100, left=120, right=120)
        p = k_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, item in enumerate(checklist_items):
        cells = t_chk.rows[row_idx + 1].cells
        cells[0].text = str(row_idx + 1)
        cells[1].text = item
        cells[2].text = "[   ] ĐẠT (PASS)      [   ] CHƯA ĐẠT (FAIL)"

        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for c in cells:
            set_cell_background(c, bg_color)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.color.rgb = COLOR_TEXT

    chk_col_widths = [Inches(0.6), Inches(4.2), Inches(2.2)]
    for row in t_chk.rows:
        for idx, width in enumerate(chk_col_widths):
            row.cells[idx].width = width

    # Save document
    output_path = r"d:\WEB\App_Dat_Xe_Tai\KE_HOACH_KIEM_THU_APP_FASTDROP.docx"
    doc.save(output_path)
    print(f"File created successfully at: {output_path}")

if __name__ == "__main__":
    create_document()
